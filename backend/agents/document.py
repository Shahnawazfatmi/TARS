"""
Document agent: turns earlier results into a Word (.docx) report.

    {"id": "s2", "tool": "create_document", "depends_on": ["s1"],
     "args": {"instruction": "A short report on AI agent trends"}}

The AI plans the content (title, sections, bullets) in one call, and
the code builds the file, so the file is always valid and looks the
same every time. Research sources become a clickable "Sources" list.
Files are saved by the shared rules in agents/files.py: only inside
the user's folders, never overwritten.
"""

import json
import logging
import os
import re
from datetime import datetime
from pathlib import Path

import config
from agents import files


logger = logging.getLogger("fatmee.document")

_llm = None

# Source numbers like [1] inside research text.
_CITE = re.compile(r"\[(\d+)\]")

# Control characters that are not allowed inside a Word file.
_BAD_XML_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

_LINK_COLOR = "0563C1"


def configure(llm) -> None:
    """Give the agent the shared LLM (called once at startup)."""

    global _llm
    _llm = llm


# ---------------- where files go ----------------

def _default_folder() -> Path:
    if config.DOCUMENTS_DIR:
        return Path(os.path.expanduser(config.DOCUMENTS_DIR))

    return files.desktop_folder("FATMEE Reports")


def resolve_folder(folder: str = "") -> Path:
    """Where to save: "" = Desktop/FATMEE Reports (see agents/files.py for the rules)."""

    return files.resolve_folder(folder, _default_folder())


# ---------------- sources ----------------

def _renumber(value, mapping: dict):
    if isinstance(value, str):
        return _CITE.sub(lambda m: f"[{mapping.get(int(m[1]), m[1])}]", value)

    if isinstance(value, list):
        return [_renumber(v, mapping) for v in value]

    if isinstance(value, dict):
        return {k: _renumber(v, mapping) for k, v in value.items()}

    return value


def collect_sources(inputs: dict) -> tuple:
    """
    Gather the sources of every research result into one numbered list.

    Each research topic numbers its own sources from [1], so the
    numbers in its text are changed to the shared list's numbers.
    The same link used twice gets one number.
    Returns (inputs with sources removed and renumbered, sources).
    """

    sources = []
    number_of = {}       # url -> shared number
    renumbered = {}

    for step_id, result in (inputs or {}).items():

        if not (isinstance(result, dict) and isinstance(result.get("sources"), list)):
            renumbered[step_id] = result
            continue

        local = {}

        for i, source in enumerate(result["sources"], start=1):

            if not isinstance(source, dict) or not source.get("url"):
                continue

            url = str(source["url"])

            if url not in number_of:
                number_of[url] = len(sources) + 1
                sources.append({
                    "n": number_of[url],
                    "title": str(source.get("title") or url),
                    "url": url,
                })

            try:
                local[int(source.get("n", i))] = number_of[url]
            except (TypeError, ValueError):
                local[i] = number_of[url]

        body = {k: v for k, v in result.items() if k != "sources"}
        renumbered[step_id] = _renumber(body, local)

    return renumbered, sources


# ---------------- content from the AI ----------------

def _clean(text) -> str:
    """Plain text: no markdown marks, no characters Word refuses."""

    text = _BAD_XML_CHARS.sub("", str(text or ""))
    text = text.replace("**", "").replace("__", "")
    text = re.sub(r"^\s*(#+|[-*•]|\d+\.)\s+", "", text)

    return text.strip()


def _as_list(value) -> list:
    if value is None:
        return []

    if not isinstance(value, list):
        value = [value]

    return [t for t in (_clean(v) for v in value) if t]


def normalize_content(raw: dict, title: str = "") -> dict:
    """Make the AI's plan safe to build: title, summary, sections."""

    sections = []

    for section in raw.get("sections") or []:

        if isinstance(section, str):
            section = {"paragraphs": [section]}

        if not isinstance(section, dict):
            continue

        heading = _clean(section.get("heading"))
        paragraphs = _as_list(section.get("paragraphs") or section.get("text"))
        bullets = _as_list(section.get("bullets") or section.get("points"))

        if heading or paragraphs or bullets:
            sections.append({"heading": heading, "paragraphs": paragraphs, "bullets": bullets})

    return {
        "title": _clean(title) or _clean(raw.get("title")) or "FATMEE report",
        "summary": _clean(raw.get("summary")),
        "sections": sections,
    }


def _plan_content(instruction: str, title: str, content: str, inputs: dict, sources: list) -> dict:

    text = json.dumps(inputs or {}, ensure_ascii=False, default=str)

    if len(text) > config.DOCUMENT_INPUT_CHARS:
        text = text[:config.DOCUMENT_INPUT_CHARS] + " ...(shortened)"

    numbered = "\n".join(f"[{s['n']}] {s['title']} - {s['url']}" for s in sources) or "(none)"

    prompt = f"""
You are FATMEE's document writer. Plan the content of a Word document.

INSTRUCTION:
{instruction or "Write a clear, well-organized report from the information below."}

TITLE (use it if given):
{title or "(choose a short title)"}

TEXT TO INCLUDE:
{content or "(none)"}

RESULTS OF EARLIER STEPS:
{text if inputs else "(none)"}

SOURCES:
{numbered}

Return JSON only, in this shape:
{{
  "title": "short title",
  "summary": "2-3 sentence introduction",
  "sections": [
    {{"heading": "...", "paragraphs": ["..."], "bullets": ["..."]}}
  ]
}}

RULES:
- Use only the information above. Do not invent facts, numbers,
  dates or sources.
- If there are no results and no text, write from general knowledge
  and leave out exact numbers or dates you are not sure of.
- After a fact that comes from a source, keep its number, like [2].
  Use only the numbers listed in SOURCES.
- 2 to 8 sections. Plain text only: no markdown, no tables.
- Do not add a sources section; FATMEE adds it.
"""

    return _llm.json(prompt, component="document")


# ---------------- building the Word file ----------------

def _add_link(paragraph, url: str, text: str) -> None:
    """A clickable link (python-docx has no built-in for this)."""

    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    rel_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)

    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), _LINK_COLOR)
    props.append(color)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)

    run.append(props)

    text_el = OxmlElement("w:t")
    text_el.text = text
    text_el.set(qn("xml:space"), "preserve")
    run.append(text_el)

    link.append(run)
    paragraph._p.append(link)


def build_docx(doc_content: dict, sources: list, path: Path) -> int:
    """Write the .docx file. Returns the number of words in it."""

    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.core_properties.title = doc_content["title"]
    doc.core_properties.author = "FATMEE"

    doc.add_heading(doc_content["title"], level=0)

    made = doc.add_paragraph().add_run(f"Made by FATMEE · {datetime.now().strftime('%d %B %Y, %H:%M')}")
    made.italic = True
    made.font.size = Pt(9)
    made.font.color.rgb = RGBColor(0x70, 0x70, 0x70)

    words = []

    if doc_content["summary"]:
        doc.add_paragraph(doc_content["summary"])
        words.append(doc_content["summary"])

    for section in doc_content["sections"]:

        if section["heading"]:
            doc.add_heading(section["heading"], level=1)
            words.append(section["heading"])

        for paragraph in section["paragraphs"]:
            doc.add_paragraph(paragraph)
            words.append(paragraph)

        for bullet in section["bullets"]:
            doc.add_paragraph(bullet, style="List Bullet")
            words.append(bullet)

    if sources:
        doc.add_heading("Sources", level=1)

        for source in sources:
            paragraph = doc.add_paragraph(f"[{source['n']}] ")
            title = _clean(source["title"]) or source["url"]

            if source["url"].startswith(("http://", "https://")):
                _add_link(paragraph, source["url"], title)
            else:
                paragraph.add_run(f"{title} - {source['url']}")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))

    return sum(len(text.split()) for text in words)


# ---------------- the tool ----------------

def create_document(
    instruction: str = "",
    title: str = "",
    content: str = "",
    file_name: str = "",
    folder: str = "",
    open_after=False,
    inputs: dict | None = None,
) -> dict:
    """Write a Word document from earlier results and save it."""

    instruction = (instruction or "").strip()
    content = (content or "").strip()

    if not (instruction or title or content or inputs):
        return {"error": "The document agent needs something to write about."}

    if _llm is None:
        return {"error": "The document agent is not configured."}

    try:
        target = resolve_folder(folder)
    except ValueError as e:
        return {"error": str(e)}

    clean_inputs, sources = collect_sources(inputs or {})

    try:
        raw = _plan_content(instruction, title, content, clean_inputs, sources)
    except Exception as e:
        logger.warning("Document content failed: %s", e)
        return {"error": "Could not write the document content."}

    doc_content = normalize_content(raw if isinstance(raw, dict) else {}, title)

    if not doc_content["sections"] and not doc_content["summary"]:
        return {"error": "The document came out empty."}

    path = files.unique_path(target, files.safe_file_name(file_name or doc_content["title"], ".docx", "FATMEE report"))

    try:
        words = build_docx(doc_content, sources, path)
    except OSError as e:
        logger.warning("Could not save %s: %s", path, e)
        return {"error": f"Could not save the document in {target} ({type(e).__name__})."}

    logger.info("Document saved: %s (%d words)", path, words)

    if files.as_bool(open_after):
        files.open_file(path)

    return {
        "saved": True,
        "path": str(path),
        "file_name": path.name,
        "folder": str(path.parent),
        "title": doc_content["title"],
        "sections": [s["heading"] for s in doc_content["sections"]],
        "words": words,
        "sources": len(sources),
    }
