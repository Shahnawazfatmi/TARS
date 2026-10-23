"""
File reader agent: reads a file on the user's PC (PDF, Word, text) and
summarizes it or answers a question about it. It can also open a file
in its usual program.

    {"id": "s1", "tool": "read_file",
     "args": {"path": "desktop/tars_research.pdf", "question": "What is it about?"}}

How a file is read:
1. The code pulls the text out (pypdf for PDF, python-docx for Word).
   The text is cached until the file changes, so a big file is only
   extracted once.
2. Short text: one AI call writes the summary and answers the question.
3. Long text: it is split into parts by page. Each part gets short
   notes (one AI call each), then one call combines the notes. If the
   file has more parts than READ_MAX_PARTS, the parts that best match
   the question are read (with no question: a spread across the file),
   and the result says how much was read.

Safety: only files inside the user's own folders (agents/files.py),
only known document types, and a size limit. Reading changes nothing.
"""

import contextvars
import difflib
import hashlib
import json
import logging
import os
import re
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import config
from agents import files


logger = logging.getLogger("fatmee.reader")

_llm = None

TEXT_TYPES = {".txt", ".md", ".csv", ".json", ".log", ".html", ".htm", ".xml", ".yaml", ".yml"}
READABLE_TYPES = {".pdf", ".docx"} | TEXT_TYPES

# open_file opens documents, pictures and media, never programs or scripts.
OPENABLE_TYPES = READABLE_TYPES | {
    ".doc", ".xlsx", ".xls", ".pptx", ".ppt", ".rtf", ".odt",
    ".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".mp4", ".mov", ".mp3", ".wav",
}

# Folders searched for a bare file name like "report.pdf".
_SEARCH_FOLDERS = ("desktop", "documents", "downloads")

# Size of one "part" of a Word or text file, which has no real pages.
_PART_CHARS = 3000

# A page with fewer characters than this has no real text (a scan, a
# picture, or only a page number).
_MIN_PAGE_CHARS = 10

_STOP_WORDS = {
    "what", "which", "about", "does", "this", "that", "with", "from", "have", "they",
    "their", "there", "where", "when", "will", "would", "could", "should", "into",
    "file", "paper", "document", "pdf", "tell", "explain", "summarize", "summary",
}


class ReadError(Exception):
    """A file that cannot be read, with a message for the user."""


def configure(llm) -> None:
    """Give the agent the shared LLM (called once at startup)."""

    global _llm
    _llm = llm


# ---------------- finding the file ----------------

def _suggest(path: Path) -> str:
    """Similar file names in the same folder, for a clear "not found"."""

    try:
        names = [p.name for p in path.parent.iterdir() if p.is_file()]
    except OSError:
        return ""

    close = difflib.get_close_matches(path.name, names, n=3, cutoff=0.5)

    return f" Did you mean: {', '.join(close)}?" if close else ""


def resolve_file(path: str) -> Path:
    """
    Turn what the user or planner gave into a real file path.

    Accepts a full path, "~/...", "desktop/name.pdf" (also documents,
    downloads, pictures: the real Windows folders, even if OneDrive
    moved them), or a bare name, which is looked for on the Desktop,
    in Documents and in Downloads. Only files inside the user's own
    folders (ValueError otherwise).
    """

    raw = (path or "").strip().strip('"').strip("'")

    if not raw:
        raise ValueError("No file given.")

    candidate = Path(os.path.expanduser(raw))

    if not candidate.is_absolute():
        parts = Path(raw.replace("\\", "/").lstrip("~/")).parts
        first = parts[0].lower() if parts else ""

        if first in files.KNOWN_FOLDERS:
            candidate = (files.known_folder(first) or Path.home()).joinpath(*parts[1:])

        else:
            candidate = Path.home() / raw

            if len(parts) == 1 and not candidate.exists():
                for name in _SEARCH_FOLDERS:
                    folder = files.known_folder(name)

                    if folder and (folder / raw).exists():
                        candidate = folder / raw
                        break

    resolved = candidate.resolve()

    if not any(resolved == root or resolved.is_relative_to(root) for root in files.allowed_roots()):
        raise ValueError(f"FATMEE only reads and opens files inside your own folders, not {resolved}.")

    if not resolved.is_file():
        raise FileNotFoundError(f"File not found: {resolved}.{_suggest(resolved)}")

    return resolved


# ---------------- pulling the text out ----------------

def _clean(text: str) -> str:
    text = (text or "").replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n\s*\n\s*\n+", "\n\n", text).strip()


def _split_parts(text: str) -> list:
    """Word and text files have no pages: split them into parts of about _PART_CHARS."""

    parts, current = [], ""

    for block in re.split(r"(\n\s*\n)", text):
        if len(current) + len(block) > _PART_CHARS and current.strip():
            parts.append(current.strip())
            current = ""
        current += block

    if current.strip():
        parts.append(current.strip())

    return [{"n": i, "text": t} for i, t in enumerate(parts, start=1)]


def _extract_pdf(path: Path) -> list:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(str(path))

        if reader.is_encrypted and not reader.decrypt(""):
            raise ReadError("This PDF is password protected, so FATMEE cannot read it.")

        pages = []

        for number, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception as e:
                # One broken page should not stop the whole file.
                logger.warning("Page %d of %s could not be read: %s", number, path.name, e)
                text = ""

            pages.append({"n": number, "text": _clean(text)})

        return pages

    except ReadError:
        raise

    except (PdfReadError, ValueError, OSError) as e:
        raise ReadError(f"This PDF looks damaged and could not be read ({type(e).__name__}).") from e


def _extract_docx(path: Path) -> list:
    from docx import Document

    try:
        document = Document(str(path))
    except Exception as e:
        raise ReadError(f"This Word file could not be opened ({type(e).__name__}).") from e

    blocks = [p.text for p in document.paragraphs if p.text.strip()]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))

    return _split_parts(_clean("\n\n".join(blocks)))


def _extract_text(path: Path) -> list:
    return _split_parts(_clean(path.read_text(encoding="utf-8", errors="replace")))


_EXTRACTORS = {".pdf": _extract_pdf, ".docx": _extract_docx}


# ---------------- cache ----------------

def _file_key(path: Path) -> str:
    stat = path.stat()
    return hashlib.sha1(f"{path}|{stat.st_mtime_ns}|{stat.st_size}".encode()).hexdigest()


def _cache_get(key: str):
    try:
        with open(os.path.join(config.READ_CACHE_DIR, f"{key}.json"), encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return None


def _cache_put(key: str, value) -> None:
    """Write atomically; a cache problem never breaks reading."""

    try:
        os.makedirs(config.READ_CACHE_DIR, exist_ok=True)
        target = os.path.join(config.READ_CACHE_DIR, f"{key}.json")
        temp = f"{target}.{os.getpid()}.tmp"

        with open(temp, "w", encoding="utf-8") as f:
            json.dump(value, f, ensure_ascii=False)

        os.replace(temp, target)

    except OSError as e:
        logger.warning("Could not write the reader cache: %s", e)


def extract_pages(path: Path) -> tuple:
    """([{"n": 1, "text": ...}, ...], cached?) for a resolved file."""

    key = _file_key(path)
    cached = _cache_get(key)

    if isinstance(cached, list):
        return cached, True

    pages = _EXTRACTORS.get(path.suffix.lower(), _extract_text)(path)
    _cache_put(key, pages)

    return pages, False


# ---------------- choosing what to read ----------------

def parse_pages(spec: str, total: int) -> set:
    """ "1-5, 9" -> {1, 2, 3, 4, 5, 9}, kept within 1..total. Empty = every page."""

    if not (spec or "").strip():
        return set(range(1, total + 1))

    chosen = set()

    for piece in re.split(r"[,;]", str(spec)):
        numbers = [int(n) for n in re.findall(r"\d+", piece)]

        if len(numbers) == 1:
            chosen.add(numbers[0])
        elif len(numbers) >= 2:
            low, high = sorted(numbers[:2])
            chosen.update(range(low, high + 1))

    return {n for n in chosen if 1 <= n <= total}


def make_chunks(pages: list, limit: int) -> list:
    """Pack pages, in order, into chunks of at most `limit` characters."""

    chunks, current = [], None

    for page in pages:
        text = page["text"]

        if not text:
            continue

        # A page longer than one chunk is split on its own.
        pieces = [text[i:i + limit] for i in range(0, len(text), limit)] or [""]

        for piece in pieces:
            block = f"[page {page['n']}]\n{piece}\n"

            if current and len(current["text"]) + len(block) > limit:
                chunks.append(current)
                current = None

            if current is None:
                current = {"first": page["n"], "last": page["n"], "text": ""}

            current["text"] += block
            current["last"] = page["n"]

    if current:
        chunks.append(current)

    return chunks


def _keywords(question: str) -> list:
    words = re.findall(r"[a-z0-9]+", (question or "").lower())
    return [w for w in words if len(w) > 3 and w not in _STOP_WORDS]


def pick_chunks(chunks: list, question: str, limit: int) -> list:
    """
    At most `limit` chunks, in file order: the ones that best match the
    question, or (no question, or nothing matches) a spread across the
    file that always includes the start.
    """

    if len(chunks) <= limit:
        return chunks

    keywords = _keywords(question)

    if keywords:
        scores = [sum(c["text"].lower().count(w) for w in keywords) for c in chunks]

        if any(scores):
            best = sorted(range(len(chunks)), key=lambda i: -scores[i])[:limit]
            return [chunks[i] for i in sorted(best)]

    spread = sorted({round(i * (len(chunks) - 1) / (limit - 1)) for i in range(limit)}) if limit > 1 else [0]

    return [chunks[i] for i in spread]


def page_ranges(numbers) -> str:
    """{1, 2, 3, 7} -> "1-3, 7"."""

    numbers = sorted(set(numbers))
    ranges, start = [], None

    for i, n in enumerate(numbers):
        if start is None:
            start = n
        if i == len(numbers) - 1 or numbers[i + 1] != n + 1:
            ranges.append(f"{start}-{n}" if n != start else str(n))
            start = None

    return ", ".join(ranges)


# ---------------- the AI part ----------------

_ANSWER_FORMAT = """Return ONLY JSON:
{
  "document_type": "what kind of file this is, e.g. research paper, invoice, CV, report",
  "summary": "4-6 sentences: what the file is about and its main message",
  "key_points": ["5-8 short points, each ending with its page, like (p. 3)"],
  "answer": "a direct answer to the user's question from the file, with pages, or \\"\\" if there is no question"
}
Use only what the file says. If it does not answer the question, say so in "answer"."""


def _ask(prompt: str) -> dict:
    return _llm.json(prompt, component="read_file")


def _notes(chunk: dict, index: int, count: int, name: str, unit: str, total: int, question: str) -> list:
    """Short notes on one part of a long file (one AI call)."""

    focus = f"what helps answer: {question}" if question else "the main ideas, findings and facts"

    result = _ask(f"""
You are FATMEE's file reader, reading part {index} of {count} of the file "{name}"
({unit}s {chunk['first']}-{chunk['last']} of {total}).

Write notes on this part, focusing on {focus}.

Return ONLY JSON: {{"notes": ["up to 8 short factual notes, each ending with its {unit}, like (p. 12)"]}}

TEXT:
{chunk['text']}
""")

    notes = result.get("notes") or []

    return [str(n) for n in notes][:8] if isinstance(notes, list) else [str(notes)]


def _summarize(chunks: list, name: str, unit: str, total: int, question: str) -> dict:
    """One call for a short file; notes per part, then one combining call, for a long one."""

    asked = f"\nUSER QUESTION:\n{question}\n" if question else "\nThe user wants a summary of the file.\n"

    if len(chunks) == 1:
        return _ask(f"""
You are FATMEE's file reader. Read the text of the file "{name}" ({total} {unit}s).
{asked}
{_ANSWER_FORMAT}

TEXT:
{chunks[0]['text']}
""")

    def notes_for(item):
        index, chunk = item
        try:
            return index, _notes(chunk, index, len(chunks), name, unit, total, question)
        except Exception as e:
            # One part failing should not lose the rest of the file.
            logger.warning("Part %d of %s could not be read: %s", index, name, e)
            return index, None

    with ThreadPoolExecutor(max_workers=max(1, config.READ_WORKERS)) as pool:
        # Each part gets its own copy of the task's context, so its AI
        # calls land in this task's log.
        futures = [
            pool.submit(contextvars.copy_context().run, notes_for, item)
            for item in enumerate(chunks, start=1)
        ]
        results = sorted(f.result() for f in futures)

    gathered = [
        f"PART {index} ({unit}s {chunks[index - 1]['first']}-{chunks[index - 1]['last']}):\n" + "\n".join(f"- {n}" for n in notes)
        for index, notes in results if notes
    ]

    if not gathered:
        raise ReadError("Could not read any part of the file.")

    return _ask(f"""
You are FATMEE's file reader. Below are NOTES FROM THE PARTS of the file "{name}"
({total} {unit}s) that were read. Combine them into one result.
{asked}
{_ANSWER_FORMAT}

NOTES:
{chr(10).join(gathered)}
""")


# ---------------- the tools ----------------

def read_file(path: str = "", question: str = "", pages: str = "") -> dict:
    """Read a file on the user's PC and summarize it or answer a question about it."""

    question = (question or "").strip()

    try:
        target = resolve_file(path)
    except (ValueError, FileNotFoundError) as e:
        return {"error": str(e)}

    suffix = target.suffix.lower()

    if suffix not in READABLE_TYPES:
        return {"error": f"FATMEE can read PDF, Word (.docx) and text files, not {suffix or 'this'} files."}

    size_mb = target.stat().st_size / (1024 * 1024)

    if size_mb > config.READ_MAX_FILE_MB:
        return {"error": f"The file is {size_mb:.0f} MB; FATMEE reads files up to {config.READ_MAX_FILE_MB:.0f} MB."}

    if _llm is None:
        return {"error": "The file reader is not configured."}

    try:
        all_pages, cached = extract_pages(target)
    except ReadError as e:
        return {"error": str(e), "file_name": target.name}
    except OSError as e:
        return {"error": f"Could not open the file ({type(e).__name__}). Is it open in another program?", "file_name": target.name}

    unit = "page" if suffix == ".pdf" else "part"
    total = len(all_pages)
    text_chars = sum(len(p["text"]) for p in all_pages)

    if suffix == ".pdf" and not any(len(p["text"]) >= _MIN_PAGE_CHARS for p in all_pages):
        return {
            "error": "This PDF has no readable text: its pages are probably scanned images. "
                     "FATMEE cannot read scanned PDFs yet (that needs OCR).",
            "file_name": target.name, "pages_total": total,
        }

    if not text_chars:
        return {"error": "The file is empty.", "file_name": target.name}

    wanted = parse_pages(pages, total)

    if not wanted:
        return {"error": f"The file has {total} {unit}s; none of '{pages}' exist.", "file_name": target.name}

    chunks = make_chunks([p for p in all_pages if p["n"] in wanted], config.READ_CHUNK_CHARS)
    chosen = pick_chunks(chunks, question, config.READ_MAX_PARTS)
    read = {p["n"] for p in all_pages if p["n"] in wanted and p["text"]
            and any(c["first"] <= p["n"] <= c["last"] for c in chosen)}
    complete = len(chosen) == len(chunks)

    # The same question on the same (unchanged) file gives the same answer.
    summary_key = hashlib.sha1(f"{_file_key(target)}|{sorted(wanted)}|{question.lower()}|{config.READ_CHUNK_CHARS}|{config.READ_MAX_PARTS}".encode()).hexdigest()
    result = _cache_get(summary_key)

    if not isinstance(result, dict):
        try:
            result = _summarize(chosen, target.name, unit, total, question)
        except ReadError as e:
            return {"error": str(e), "file_name": target.name}
        except ValueError:
            return {"error": "Could not summarize the file.", "file_name": target.name}

        _cache_put(summary_key, result)
    else:
        cached = True

    points = result.get("key_points") or []

    reply = {
        "file_name": target.name,
        "path": str(target),
        "file_type": suffix.lstrip("."),
        "size_mb": round(size_mb, 1),
        "unit": unit,
        "total": total,
        "read": len(read),
        "pages_read": page_ranges(read),
        "complete": complete,
        "document_type": str(result.get("document_type", "")),
        "summary": str(result.get("summary", "")),
        "key_points": [str(p) for p in points][:8] if isinstance(points, list) else [str(points)],
        "cached": cached,
    }

    if question:
        reply["question"] = question
        reply["answer"] = str(result.get("answer", ""))

    if not complete:
        how = "the parts that best match the question" if question and _keywords(question) else "a spread across the file"
        reply["note"] = (f"Long file: read {len(read)} of {total} {unit}s ({how}). "
                         f"Ask about a topic or give pages (e.g. pages '40-60') to read other parts.")

    empty = sum(1 for p in all_pages if p["n"] in wanted and not p["text"])

    if suffix == ".pdf" and empty > len(wanted) / 3:
        reply["warning"] = f"{empty} of {len(wanted)} pages have no text (probably pictures or scans) and were skipped."

    return reply


def open_file(path: str = "") -> dict:
    """Open a file in its usual program (PDF viewer, Word, photo viewer...)."""

    try:
        target = resolve_file(path)
    except (ValueError, FileNotFoundError) as e:
        return {"error": str(e)}

    if target.suffix.lower() not in OPENABLE_TYPES:
        return {"error": f"FATMEE opens documents, pictures and media only, not {target.suffix or 'this'} files."}

    if os.name != "nt":
        return {"error": "Opening files works on Windows only."}

    try:
        os.startfile(target)
    except OSError as e:
        return {"error": f"Windows could not open the file ({e.strerror or type(e).__name__})."}

    return {"opened": True, "file_name": target.name, "path": str(target)}
