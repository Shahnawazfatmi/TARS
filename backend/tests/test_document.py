import json
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

import config
import main
from agents import document
from conftest import FakeProvider
from executor import execute_plan
from llm import LLM
from tools import registry


CONTENT = json.dumps({
    "title": "AI Agent Trends",
    "summary": "Agents are moving into **real work** [1].",
    "sections": [
        {"heading": "## What changed", "paragraphs": ["Tools got better [3]."], "bullets": ["- Memory", "* Planning"]},
        "A section that is just text.",
        {"heading": "", "paragraphs": [], "bullets": []},
    ],
})

RESEARCH_A = {
    "topic": "agents", "summary": "Agents grow [1].", "key_points": ["Planning [2]"],
    "sources": [{"n": 1, "title": "A1", "url": "https://a.example/1"},
                {"n": 2, "title": "Shared", "url": "https://shared.example"}],
}
RESEARCH_B = {
    "topic": "rules", "summary": "New rules [1] and [2].",
    "sources": [{"n": 1, "title": "B1", "url": "https://b.example/1"},
                {"n": 2, "title": "Shared", "url": "https://shared.example"}],
}


@pytest.fixture
def docs_dir(tmp_path, monkeypatch):
    """Save test documents in a temp folder, never on the real Desktop."""

    monkeypatch.setattr(config, "DOCUMENTS_DIR", str(tmp_path))
    monkeypatch.setattr(document, "_allowed_roots", lambda: [tmp_path.resolve()])
    return tmp_path


@pytest.fixture
def fake_writer(monkeypatch):
    fake = FakeProvider({"document writer": CONTENT})
    monkeypatch.setattr(document, "_llm", LLM([fake]))
    return fake


# ---------------- file names and folders ----------------

@pytest.mark.parametrize("name,expected", [
    ("Q3: AI / trends?", "Q3 AI trends.docx"),
    ("report.docx", "report.docx"),
    ("  ...  ", "FATMEE report.docx"),
    ("", "FATMEE report.docx"),
])
def test_safe_file_name(name, expected):
    assert document.safe_file_name(name) == expected


def test_never_overwrites(tmp_path):
    (tmp_path / "r.docx").write_text("x")
    (tmp_path / "r (2).docx").write_text("x")
    assert document.unique_path(tmp_path, "r.docx").name == "r (3).docx"


def test_folder_outside_user_folders_is_refused(docs_dir):
    with pytest.raises(ValueError, match="only saves documents"):
        document.resolve_folder("C:/Windows/System32")

    assert document.resolve_folder("") == docs_dir.resolve()
    assert document.resolve_folder(str(docs_dir / "sub")) == (docs_dir / "sub").resolve()


def test_desktop_is_found():
    assert document._known_folder("desktop").is_dir()


# ---------------- sources ----------------

def test_sources_from_several_topics_share_one_list():
    inputs, sources = document.collect_sources({"s1": RESEARCH_A, "s2": RESEARCH_B, "s3": {"output": "x [1]"}})

    assert [s["url"] for s in sources] == ["https://a.example/1", "https://shared.example", "https://b.example/1"]
    # topic B's [1] is now [3]; its [2] is the shared link, [2]
    assert inputs["s2"]["summary"] == "New rules [3] and [2]."
    assert inputs["s1"]["key_points"] == ["Planning [2]"]
    assert "sources" not in inputs["s1"]
    # results without sources are untouched
    assert inputs["s3"] == {"output": "x [1]"}


def test_content_is_cleaned():
    content = document.normalize_content(json.loads(CONTENT))
    assert content["summary"] == "Agents are moving into real work [1]."
    assert content["sections"][0] == {"heading": "What changed", "paragraphs": ["Tools got better [3]."], "bullets": ["Memory", "Planning"]}
    assert content["sections"][1]["paragraphs"] == ["A section that is just text."]
    # the empty section is dropped
    assert len(content["sections"]) == 2


# ---------------- the tool ----------------

def test_create_document_builds_a_real_word_file(docs_dir, fake_writer):
    result = document.create_document(
        instruction="Report on AI trends", inputs={"s1": RESEARCH_A, "s2": RESEARCH_B},
    )

    path = Path(result["path"])
    assert result["saved"] and path.exists() and path.parent == docs_dir
    assert result["file_name"] == "AI Agent Trends.docx"
    assert result["sources"] == 3 and result["words"] > 5

    # the writer saw renumbered research and the shared source list
    prompt = fake_writer.calls[0]["prompt"]
    assert "New rules [3] and [2]." in prompt and "[3] B1 - https://b.example/1" in prompt

    doc = Document(str(path))
    texts = [p.text for p in doc.paragraphs]
    styles = [p.style.name for p in doc.paragraphs]
    assert texts[0] == "AI Agent Trends" and styles[0] == "Title"
    assert "What changed" in texts and "Sources" in texts
    assert styles[texts.index("Memory")] == "List Bullet"
    # sources are clickable links
    links = [r.target_ref for r in doc.part.rels.values() if "hyperlink" in r.reltype]
    assert links == ["https://a.example/1", "https://shared.example", "https://b.example/1"]
    assert any(t.startswith("[3] ") and "B1" in t for t in texts)


def test_same_title_twice_makes_two_files(docs_dir, fake_writer):
    first = document.create_document(instruction="x")
    second = document.create_document(instruction="x")
    assert first["file_name"] == "AI Agent Trends.docx"
    assert second["file_name"] == "AI Agent Trends (2).docx"


def test_title_and_file_name_from_user_win(docs_dir, fake_writer):
    result = document.create_document(title="My Notes", file_name="notes", content="hello")
    assert result["title"] == "My Notes" and result["file_name"] == "notes.docx"
    assert "hello" in fake_writer.calls[0]["prompt"]


def test_document_errors(docs_dir, monkeypatch):
    assert "needs something" in document.create_document()["error"]

    monkeypatch.setattr(document, "_llm", LLM([FakeProvider({"document writer": "not json"})]))
    assert "Could not write" in document.create_document(instruction="x")["error"]

    monkeypatch.setattr(document, "_llm", LLM([FakeProvider({"document writer": '{"sections": []}'})]))
    assert "empty" in document.create_document(instruction="x")["error"]

    assert "only saves" in document.create_document(instruction="x", folder="C:/Windows")["error"]


# ---------------- in a plan ----------------

def test_executor_gives_the_document_its_inputs(docs_dir, fake_writer):
    results = execute_plan({"steps": [
        {"id": "s1", "tool": "get_ram_usage"},
        {"id": "s2", "tool": "create_document", "depends_on": ["s1"],
         "args": {"instruction": "RAM report"}},
    ]})

    assert results["create_document"]["saved"]
    prompt = fake_writer.calls[0]["prompt"]
    assert '"s1"' in prompt and "total" in prompt
    assert registry.takes_inputs("create_document") and not registry.takes_inputs("get_ram_usage")


def test_api_research_style_plan_makes_document(docs_dir, fake_llm, monkeypatch):
    plan = json.dumps({"goal": "ram report", "steps": [
        {"id": "s1", "step": 1, "action": "ram", "tool": "get_ram_usage", "args": {}, "depends_on": []},
        {"id": "s2", "step": 2, "action": "doc", "tool": "create_document",
         "args": {"instruction": "RAM report"}, "depends_on": ["s1"]},
    ]})
    fake = fake_llm({
        "You are the planning component": plan,
        "document writer": CONTENT,
        "Goal Verification": '{"goal_satisfied": true, "reason": "ok", "missing_information": []}',
    })
    monkeypatch.setattr(document, "_llm", main.llm)

    body = TestClient(main.app).post("/ask", json={"question": "Make a Word report of my RAM"}).json()
    assert body["tools_used"] == ["get_ram_usage", "create_document"]
    assert (docs_dir / "AI Agent Trends.docx").exists()
    # planner was taught to use create_document without an extra think step
    assert "create_document" in fake.calls[0]["prompt"]
    # the final answer is told to say where the file is
    final = [c for c in fake.calls if c["system"] and "document was saved" in c["system"]]
    assert final
