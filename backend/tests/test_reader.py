import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

import cards
import config
import main
from agents import files, reader
from conftest import FakeProvider
from llm import LLM
from tools.registry import TOOLS_BY_NAME


SUMMARY = json.dumps({
    "document_type": "research paper",
    "summary": "A paper about TARS, a robot with adjustable humor.",
    "key_points": ["TARS has a humor setting (p. 2)"],
    "answer": "",
})
NOTES = json.dumps({"notes": ["A fact (p. 1)"]})


def make_pdf(path: Path, pages: list) -> Path:
    """A small real PDF with one line of text per page ("" = a page with no text)."""

    objects = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []

    for text in pages:
        safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({safe}) Tj ET" if text else ""
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                       f"/Resources << /Font << /F1 3 0 R >> >> /Contents {len(objects)} 0 R >>")
        kids.append(f"{len(objects)} 0 R")

    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"

    out, offsets = b"%PDF-1.4\n", []

    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1")

    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()

    path.write_bytes(out)
    return path


@pytest.fixture
def home(tmp_path, monkeypatch):
    """A fake home with Desktop, Documents and Downloads; FATMEE may only read inside it."""

    for name in ("Desktop", "Documents", "Downloads", "Pictures"):
        (tmp_path / name).mkdir()

    monkeypatch.setattr(files, "allowed_roots", lambda: [tmp_path.resolve()])
    monkeypatch.setattr(files, "known_folder", lambda name: tmp_path / name.capitalize())
    return tmp_path


@pytest.fixture
def model(monkeypatch):
    """The reader's AI: scripted replies, every call recorded."""

    def install(replies):
        fake = FakeProvider(replies)
        reader.configure(LLM([fake]))
        return fake

    yield install
    reader.configure(main.llm)


def fresh_replies(single=SUMMARY, notes=NOTES, combined=SUMMARY):
    # Most specific marker first: the combining prompt also says "file reader".
    return {"NOTES FROM THE PARTS": combined, "reading part": notes, "Read the text of the file": single}


# ---------------- finding the file ----------------

def test_finds_files_the_way_people_name_them(home):
    pdf = make_pdf(home / "Desktop" / "tars.pdf", ["hello"])
    doc = home / "Documents" / "notes.txt"
    doc.write_text("x")

    assert reader.resolve_file("desktop/tars.pdf") == pdf.resolve()
    assert reader.resolve_file(str(pdf)) == pdf.resolve()
    # a bare name is looked for on the Desktop, in Documents and Downloads
    assert reader.resolve_file("notes.txt") == doc.resolve()


def test_refuses_files_outside_your_folders_and_suggests_names(home):
    make_pdf(home / "Desktop" / "tars_research.pdf", ["hello"])

    with pytest.raises(ValueError, match="only reads and opens files inside your own folders"):
        reader.resolve_file("C:/Windows/win.ini")

    with pytest.raises(FileNotFoundError, match="Did you mean: tars_research.pdf"):
        reader.resolve_file("desktop/tars_reserch.pdf")


# ---------------- reading ----------------

def test_short_file_is_one_ai_call_with_the_answer(home, model):
    fake = model(fresh_replies(single=json.dumps({
        "document_type": "notes", "summary": "Meeting notes.", "key_points": ["Ship Friday (p. 1)"],
        "answer": "They ship on Friday.",
    })))
    (home / "Desktop" / "notes.txt").write_text("We agreed to ship on Friday.\n\nBudget is fine.")

    result = reader.read_file("desktop/notes.txt", question="When do they ship?")

    assert len(fake.calls) == 1 and fake.calls[0]["json_mode"]
    assert "We agreed to ship on Friday" in fake.calls[0]["prompt"]
    assert result["answer"] == "They ship on Friday." and result["complete"] is True
    assert result["read"] == result["total"] == 1 and result["unit"] == "part"


def test_pdf_pages_and_page_choice(home, model):
    fake = model(fresh_replies())
    make_pdf(home / "Desktop" / "paper.pdf", ["Intro page", "Method page", "Results page", "End page"])

    result = reader.read_file("desktop/paper.pdf", pages="2-3")

    prompt = fake.calls[0]["prompt"]
    assert "[page 2]" in prompt and "[page 3]" in prompt and "Intro page" not in prompt
    assert (result["total"], result["read"], result["pages_read"]) == (4, 2, "2-3")
    assert result["summary"].startswith("A paper about TARS")

    assert "none of '9' exist" in reader.read_file("desktop/paper.pdf", pages="9")["error"]


def test_page_helpers():
    assert reader.parse_pages("1-3, 7; 9-8", 10) == {1, 2, 3, 7, 8, 9}
    assert reader.parse_pages("", 3) == {1, 2, 3}
    assert reader.parse_pages("50", 3) == set()
    assert reader.page_ranges({1, 2, 3, 7, 9, 10}) == "1-3, 7, 9-10"


def test_long_file_is_read_in_parts_then_combined(home, model, monkeypatch):
    monkeypatch.setattr(config, "READ_CHUNK_CHARS", 120)
    monkeypatch.setattr(config, "READ_MAX_PARTS", 10)
    fake = model(fresh_replies())
    make_pdf(home / "Desktop" / "long.pdf", [f"Page {n} " + "word " * 15 for n in range(1, 7)])

    result = reader.read_file("desktop/long.pdf")

    parts = [c for c in fake.calls if "reading part" in c["prompt"]]
    assert len(parts) > 1 and len(fake.calls) == len(parts) + 1
    assert "NOTES FROM THE PARTS" in fake.calls[-1]["prompt"] and "A fact (p. 1)" in fake.calls[-1]["prompt"]
    assert result["complete"] is True and result["read"] == 6 and "note" not in result


def test_very_long_file_reads_the_parts_that_match_the_question(home, model, monkeypatch):
    monkeypatch.setattr(config, "READ_CHUNK_CHARS", 60)
    monkeypatch.setattr(config, "READ_MAX_PARTS", 2)
    fake = model(fresh_replies())
    texts = [f"Page {n} about engines and wheels" for n in range(1, 9)]
    texts[5] = "Page 6 about TARS humor settings"
    make_pdf(home / "Desktop" / "big.pdf", texts)

    result = reader.read_file("desktop/big.pdf", question="What is the TARS humor setting?")

    read_prompts = " ".join(c["prompt"] for c in fake.calls if "reading part" in c["prompt"])
    assert "TARS humor settings" in read_prompts
    assert result["complete"] is False and "best match the question" in result["note"]

    # with no question: a spread that always starts at the beginning
    chunks = reader.make_chunks([{"n": n, "text": f"Page {n} text"} for n in range(1, 21)], 20)
    assert reader.pick_chunks(chunks, "", 3)[0]["first"] == 1


def test_one_failing_part_does_not_lose_the_file(home, model, monkeypatch):
    monkeypatch.setattr(config, "READ_CHUNK_CHARS", 60)
    model(fresh_replies(notes=[NOTES, "not json", NOTES, NOTES, NOTES, NOTES, NOTES, NOTES]))
    make_pdf(home / "Desktop" / "parts.pdf", [f"Page {n} " + "text " * 8 for n in range(1, 5)])

    assert reader.read_file("desktop/parts.pdf")["summary"].startswith("A paper")


def test_scanned_pdf_says_it_needs_ocr(home, model):
    fake = model(fresh_replies())
    make_pdf(home / "Desktop" / "scan.pdf", ["", "", ""])

    result = reader.read_file("desktop/scan.pdf")

    assert "scanned images" in result["error"] and "OCR" in result["error"]
    assert fake.calls == []


def test_word_files_are_read(home, model):
    from docx import Document

    fake = model(fresh_replies())
    document = Document()
    document.add_paragraph("FATMEE quarterly plan")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Budget", "10k"
    document.save(home / "Documents" / "plan.docx")

    reader.read_file("documents/plan.docx")

    assert "FATMEE quarterly plan" in fake.calls[0]["prompt"] and "Budget | 10k" in fake.calls[0]["prompt"]


def test_same_file_is_read_once_until_it_changes(home, model):
    fake = model(fresh_replies())
    note = home / "Desktop" / "memo.txt"
    note.write_text("First version.")

    assert reader.read_file("desktop/memo.txt")["cached"] is False
    again = reader.read_file("desktop/memo.txt")
    assert again["cached"] is True and len(fake.calls) == 1

    note.write_text("Second, longer version.")
    reader.read_file("desktop/memo.txt")
    assert len(fake.calls) == 2 and "Second, longer version." in fake.calls[1]["prompt"]


def test_unsupported_too_big_and_missing(home, model, monkeypatch):
    model(fresh_replies())
    (home / "Desktop" / "setup.exe").write_bytes(b"MZ")
    (home / "Desktop" / "big.txt").write_text("x" * 2048)

    assert "not .exe files" in reader.read_file("desktop/setup.exe")["error"]
    monkeypatch.setattr(config, "READ_MAX_FILE_MB", 0.001)
    assert "reads files up to" in reader.read_file("desktop/big.txt")["error"]
    assert "File not found" in reader.read_file("desktop/nothing.pdf")["error"]


# ---------------- opening ----------------

def test_open_file_opens_documents_never_programs(home, monkeypatch):
    opened = []
    monkeypatch.setattr(reader.os, "startfile", opened.append, raising=False)
    monkeypatch.setattr(reader.os, "name", "nt")
    pdf = make_pdf(home / "Desktop" / "tars.pdf", ["hi"])
    (home / "Desktop" / "run.bat").write_text("echo hi")

    assert reader.open_file("desktop/tars.pdf")["opened"] is True and opened == [pdf.resolve()]
    assert "not .bat files" in reader.open_file("desktop/run.bat")["error"]
    assert "inside your own folders" in reader.open_file("C:/Windows/notepad.exe")["error"]
    assert len(opened) == 1


# ---------------- screen and planner ----------------

def test_card_and_screen_can_open_the_read_file(home, model):
    model(fresh_replies())
    pdf = make_pdf(home / "Desktop" / "tars.pdf", ["TARS humor page " * 3])

    card = cards.card("read_file", reader.read_file("desktop/tars.pdf"))

    assert card["type"] == "file" and card["coverage"] == "Read 1 of 1 pages" and card["title"] == "tars.pdf"
    assert cards.servable(str(pdf)) == pdf.resolve()


def test_planner_sees_both_tools():
    assert "read_file" in TOOLS_BY_NAME and "open_file" in TOOLS_BY_NAME
    assert not TOOLS_BY_NAME["open_file"].needs_approval


def test_read_a_file_through_the_api(home, fake_llm, monkeypatch):
    make_pdf(home / "Desktop" / "tars.pdf", ["TARS is a robot with a humor setting of 75 percent."])
    plan = json.dumps({"goal": "summarize", "steps": [
        {"id": "s1", "tool": "read_file", "args": {"path": "desktop/tars.pdf"}, "depends_on": []},
    ]})
    fake_llm({
        "You are the planning component": plan,
        "Read the text of the file": SUMMARY,
        "Goal Verification": '{"goal_satisfied": true, "reason": "ok", "missing_information": []}',
    })
    monkeypatch.setattr(reader, "_llm", main.llm)

    reply = TestClient(main.app).post("/ask", json={"question": "Summarize tars.pdf on my desktop"}).json()

    assert reply["tools_used"] == ["read_file"]
    assert reply["cards"][0]["type"] == "file"
